"""
多格式文本抽取服务.

从上传文件 (PDF/DOCX/TXT/CSV) 与网页 HTML 中抽取纯文本，
供"智能分析"的 Agent 链路上游使用。
"""

from __future__ import annotations

import base64
import io
from html.parser import HTMLParser

MAX_CHARS = 100_000


class ServiceNotReadyError(RuntimeError):
    """抽取依赖未安装 (镜像重建前调用会抛出, 由路由转为 503)."""


# ─── PDF ───────────────────────────────────────────────────────


def extract_pdf(data: bytes) -> tuple[str, bool]:
    """
    抽取 PDF 文本 (逐页混合): 有文字层的页直取文字, 空白页自动渲染为图片并 OCR.

    Returns:
        (text, is_scanned): is_scanned=True 表示整份 PDF 无任何文字层 (全扫描件);
        混排 PDF (部分页有文字层) 返回 False, 但仍对无文字层的页逐页 OCR。
    """
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ServiceNotReadyError("PDF 抽取组件未就绪, 请稍后重试") from e

    reader = PdfReader(io.BytesIO(data))
    page_texts: list[str] = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        page_texts.append(text)

    # 无文字层 (需 OCR) 的页索引
    blank_idx = [i for i, t in enumerate(page_texts) if not t.strip()]
    if not blank_idx:
        return "\n".join(page_texts), False

    # 逐页混合: 仅对空白页渲染位图走 OCR, 有文字层的页保留原文
    ocr_map = _ocr_pdf_pages(data, blank_idx)
    merged = [
        t.strip() if t.strip() else ocr_map.get(i, "")
        for i, t in enumerate(page_texts)
    ]
    is_scanned = len(blank_idx) == len(page_texts)  # 全部页都无文字层 → 整份扫描件
    return "\n\n".join(seg for seg in merged if seg), is_scanned


def _ocr_pdf_pages(data: bytes, page_indices: list[int], dpi: int = 200) -> dict[int, str]:
    """将指定页索引渲染为 PNG 并 OCR, 返回 {页索引: 识别文本}.

    依赖 PyMuPDF (fitz) 渲染页面; 渲染/OCR 不可用时抛 ServiceNotReadyError 转 503。
    """
    try:
        import fitz  # PyMuPDF
    except ImportError as e:
        raise ServiceNotReadyError(
            "扫描版 PDF 需要 OCR 组件 (pymupdf), 请稍后重试"
        ) from e

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as e:
        raise ServiceNotReadyError(f"扫描版 PDF 解析失败: {e}") from e

    ocr_map: dict[int, str] = {}
    try:
        zoom = dpi / 72  # PDF 基准 72dpi, 提高采样率以改善小字识别
        matrix = fitz.Matrix(zoom, zoom)
        total = doc.page_count
        for idx in page_indices:
            if idx >= total:
                continue
            pixmap = doc[idx].get_pixmap(matrix=matrix)
            ocr_text, _ = extract_image(pixmap.tobytes("png"))
            if ocr_text.strip():
                ocr_map[idx] = ocr_text.strip()
    finally:
        doc.close()

    return ocr_map


# ─── DOCX ──────────────────────────────────────────────────────


def extract_docx(data: bytes) -> tuple[str, bool]:
    """抽取 Word (.docx) 段落与表格文本 (忽略页眉页脚)."""
    try:
        from docx import Document
    except ImportError as e:
        raise ServiceNotReadyError("DOCX 抽取组件未就绪, 请稍后重试") from e

    doc = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), False


# ─── TXT / CSV / MARKDOWN ──────────────────────────────────────


def extract_txt_csv(data: bytes) -> tuple[str, bool]:
    """按 utf-8-sig → gb18030 → latin-1 依次尝试解码 (md 同为纯文本)."""
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding), False
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), False


# ─── RTF ───────────────────────────────────────────────────────


# 出现即跳过的"属性组"开头 (紧跟 { )
_RTF_SKIP_GROUP_PREFIXES = (
    "\\*", "\\fonttbl", "\\colortbl", "\\stylesheet", "\\info",
    "\\pict", "\\object", "\\generator", "\\header", "\\footer",
    "\\listtable", "\\listoverridetable", "\\revtbl", "\\rsidtbl",
)


def extract_rtf(data: bytes) -> tuple[str, bool]:
    """
    纯标准库抽取 RTF 文本 (单遍扫描 + 分组栈).

    RTF 是 ANSI 字节流, 先按 latin-1 无损解码再剥离控制字,
    收尾再尝试还原为 utf-8 / gb18030 / cp1252。
    """
    raw = data.decode("latin-1", errors="replace")
    text_parts: list[str] = []
    depth = 0
    skip_depth = -1  # >=0 表示正处在应跳过的属性组内 (记录其深度)

    def _flush_skip(d: int) -> int:
        if skip_depth >= 0 and d <= skip_depth:
            return -1
        return skip_depth

    i, n = 0, len(raw)
    while i < n:
        ch = raw[i]

        if ch == "{":
            if skip_depth < 0:
                rest = raw[i + 1:i + 40].lstrip()
                if rest.startswith(_RTF_SKIP_GROUP_PREFIXES):
                    skip_depth = depth  # 本组一闭合即恢复输出
            depth += 1
            i += 1
            continue

        if ch == "}":
            depth = max(0, depth - 1)
            skip_depth = _flush_skip(depth)
            i += 1
            continue

        if ch == "\\":
            nxt = raw[i + 1:i + 2]
            if nxt == "'":  # \'hh 十六进制字节
                if skip_depth < 0:
                    try:
                        text_parts.append(chr(int(raw[i + 2:i + 4], 16)))
                    except ValueError:
                        pass
                i += 4
                continue
            j = i + 1
            while j < n and (raw[j].isalpha() or raw[j] == "*"):
                j += 1
            word = raw[i + 1:j]
            if not word:  # 转义字面量 \{ \} \\
                if j < n and raw[j] in "{}":
                    if skip_depth < 0:
                        text_parts.append(raw[j])
                i = j + 1
                continue
            # 控制字数字参数与尾随空格
            m = j
            sign = 1
            if m < n and raw[m] == "-":
                sign = -1
                m += 1
            num = 0
            while m < n and raw[m].isdigit():
                num = num * 10 + int(raw[m])
                m += 1
            if m < n and raw[m] == " ":
                m += 1

            if skip_depth < 0:
                if word in ("par", "line"):
                    text_parts.append("\n")
                elif word == "tab":
                    text_parts.append(" ")
                elif word == "u":
                    text_parts.append(chr((sign * num) & 0xFFFF))
                    # 吞掉 \uc 规定的回退内容 (最多两个 \'hh 或单个 '?')
                    for _ in range(2):
                        if raw[m:m + 1] == "\\" and raw[m + 1:m + 2] == "'":
                            m += 4
                        else:
                            break
                    if m < n and raw[m] == "?":
                        m += 1
                # 其余控制字 (字体/字号/颜色等) 一律丢弃
            i = m
            continue

        if ch in "\r\n":
            i += 1
            continue

        if skip_depth < 0:
            text_parts.append(ch)
        i += 1

    text = "".join(text_parts)
    # 还原非 latin-1 编码的中文 (ANSI/GB2312/UTF-8 三种常见形态)
    for enc in ("utf-8", "gb18030", "cp1252"):
        try:
            text = text.encode("latin-1").decode(enc)
            break
        except (UnicodeEncodeError, UnicodeDecodeError):
            continue
    lines = [ln.strip() for ln in text.split("\n")]
    return "\n".join(ln for ln in lines if ln), False


# ─── HTML ──────────────────────────────────────────────────────


class _HtmlToText(HTMLParser):
    """标准库 HTMLParser 实现, 只保留正文语义块与标题, 不引入 bs4."""

    _BLOCK_TAGS = {
        "p", "div", "section", "article", "blockquote", "pre",
        "li", "td", "th", "h1", "h2", "h3", "h4", "h5", "h6",
        "br", "tr", "ul", "ol", "table", "header", "footer",
    }
    _SKIP_TAGS = {"script", "style", "noscript", "svg", "template"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title: str = ""
        self._buf: list[str] = []
        self._in_title = False
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP_TAGS:
            self._skip_depth += 1
        elif tag == "title":
            self._in_title = True
        elif tag in self._BLOCK_TAGS and not self._skip_depth:
            self._flush_buf()

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP_TAGS and self._skip_depth:
            self._skip_depth -= 1
        elif tag == "title":
            self._in_title = False
        elif tag in self._BLOCK_TAGS and not self._skip_depth:
            self._flush_buf()

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data
        elif not self._skip_depth and data.strip():
            self._buf.append(data.strip())

    def _flush_buf(self) -> None:
        if self._buf:
            self.parts.append(" ".join(self._buf))
            self._buf = []

    def close(self) -> None:
        super().close()
        self._flush_buf()


def html_to_text(html: str) -> tuple[str, str]:
    """HTML → (title, text)."""
    parser = _HtmlToText()
    try:
        parser.feed(html)
        parser.close()
    except Exception:
        pass
    title = " ".join(parser.title.split())[:200]
    text = "\n".join(p for p in parser.parts if p)
    return title, text


def extract_html(data: bytes) -> tuple[str, bool]:
    """抽取 .html/.htm 文件文本, 页面标题作为首行保留."""
    raw = data.decode("utf-8", errors="replace")
    title, text = html_to_text(raw)
    if title:
        text = f"{title}\n\n{text}"
    return text, False


# ─── XLSX ──────────────────────────────────────────────────────


def extract_xlsx(data: bytes) -> tuple[str, bool]:
    """抽取 Excel (.xlsx) 各工作表文本 (只读模式, 取公式缓存值)."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ServiceNotReadyError("XLSX 抽取组件未就绪, 请稍后重试") from e

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            if len(wb.worksheets) > 1:
                parts.append(f"=== {ws.title} ===")
            for row in ws.iter_rows(values_only=True):
                cells = []
                for value in row:
                    if value is None:
                        continue
                    if hasattr(value, "strftime"):
                        value = value.strftime("%Y-%m-%d %H:%M:%S")
                    elif isinstance(value, bool):
                        value = "是" if value else "否"
                    elif isinstance(value, float) and value.is_integer():
                        value = str(int(value))
                    cells.append(str(value).strip())
                if cells:
                    parts.append(" | ".join(c for c in cells if c))
    finally:
        wb.close()
    return "\n".join(parts), False


# ─── PPTX ──────────────────────────────────────────────────────


def _collect_pptx_shapes(shapes) -> list[str]:
    """递归收集幻灯片中的文本 (文本框/组合/表格)."""
    collected: list[str] = []
    for shape in shapes:
        if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
            collected.append(shape.text_frame.text.strip())
        inner_shapes = getattr(shape, "shapes", None)
        if inner_shapes is not None:
            collected.extend(_collect_pptx_shapes(inner_shapes))
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    collected.append(" | ".join(cells))
    return collected


def extract_pptx(data: bytes) -> tuple[str, bool]:
    """抽取 PowerPoint (.pptx) 各页文本."""
    try:
        from pptx import Presentation
    except ImportError as e:
        raise ServiceNotReadyError("PPTX 抽取组件未就绪, 请稍后重试") from e

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        page_parts = _collect_pptx_shapes(slide.shapes)
        if page_parts:
            if len(prs.slides) > 1:
                parts.append(f"=== 第 {idx} 页 ===")
            parts.extend(p for p in page_parts if p)
    return "\n".join(parts), False


# ─── IMAGE (OCR / 多模态 LLM) ──────────────────────────────


def _guess_image_mime(data: bytes) -> str:
    """按文件头推断图片 MIME (无需第三方库)."""
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png"
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "image/gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    if data[:2] == b"BM":
        return "image/bmp"
    return "image/png"  # 兜底, 多模态模型大多兼容


def _extract_image_tesseract(data: bytes) -> tuple[str, bool]:
    """
    本地 Tesseract OCR 抽取文本 (合同拍照/扫描件).

    依赖 Pillow + pytesseract + 系统 Tesseract (含 chi_sim/eng 语言包)。
    Tesseract 未安装或图片解码失败时抛 ServiceNotReadyError, 由路由转 503。
    """
    try:
        from PIL import Image
        import pytesseract
    except ImportError as e:
        raise ServiceNotReadyError("图片 OCR 组件未就绪, 请稍后重试") from e

    try:
        img = Image.open(io.BytesIO(data))
    except Exception as e:
        raise ServiceNotReadyError(f"图片解码失败: {e}") from e

    # 转灰度提升小字识别率; Tesseract 缺失时 image_to_string 抛 TesseractNotFoundError
    try:
        text = pytesseract.image_to_string(img.convert("L"), lang="chi_sim+eng")
    except Exception as e:
        raise ServiceNotReadyError(f"OCR 引擎不可用: {e}") from e

    return text, True


def extract_image_via_llm(data: bytes) -> tuple[str, bool]:
    """
    多模态 LLM 图片文字抽取 (高准确率, 支持中文/表格/手写).

    复用 OpenAI 兼容视觉接口 (默认通义千问 VL); 未配置视觉密钥时抛
    ServiceNotReadyError 以便上层回退到 Tesseract。
    """
    try:
        from langchain_openai import ChatOpenAI  # noqa: F401  (确保依赖可用)
    except ImportError as e:
        raise ServiceNotReadyError("多模态 LLM 组件未就绪, 请稍后重试") from e

    from app.agent.llm import get_vision_llm
    from app.core.config import get_settings

    settings = get_settings()
    if not settings.llm_vision_api_key:
        raise ServiceNotReadyError("未配置视觉模型 API Key, 回退本地 OCR")

    mime = _guess_image_mime(data)
    data_uri = f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"

    llm = get_vision_llm()
    messages = [
        (
            "user",
            [
                {
                    "type": "text",
                    "text": (
                        "请完整提取图片中的全部文字内容, 保留原有段落、标题与表格排版, "
                        "只输出文字本身, 不要任何解释或前缀。"
                    ),
                },
                {"type": "image_url", "image_url": {"url": data_uri}},
            ],
        )
    ]

    try:
        resp = llm.invoke(messages)
    except Exception as e:  # 网络/鉴权/配额等, 交由上层决定回退
        raise ServiceNotReadyError(f"多模态 LLM 调用失败: {e}") from e

    content = getattr(resp, "content", None) or ""
    if isinstance(content, list):  # 部分 provider 返回 list[block]
        content = "".join(b.get("text", "") for b in content if isinstance(b, dict))
    return content, True


def extract_image(data: bytes) -> tuple[str, bool]:
    """
    图片文字抽取统一入口: 多模态 LLM 为主, Tesseract OCR 兜底。

    策略由 settings.image_ocr_strategy 控制:
      - auto     : 配置了视觉模型则优先 LLM, 失败/未配置回退 Tesseract (默认)
      - llm      : 仅用 LLM, 不可用即抛 503
      - tesseract: 仅用 Tesseract OCR
    """
    from app.core.config import get_settings

    strategy = get_settings().image_ocr_strategy.lower()
    if strategy == "tesseract":
        return _extract_image_tesseract(data)
    if strategy == "llm":
        return extract_image_via_llm(data)

    # auto: 优先多模态 LLM, 任意失败回退本地 OCR
    try:
        return extract_image_via_llm(data)
    except ServiceNotReadyError:
        return _extract_image_tesseract(data)


# ─── 统一入口 ──────────────────────────────────────────────────


_EXTRACTORS: dict[str, object] = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "txt": extract_txt_csv,
    "csv": extract_txt_csv,
    "md": extract_txt_csv,
    "html": extract_html,
    "htm": extract_html,
    "rtf": extract_rtf,
    "xlsx": extract_xlsx,
    "pptx": extract_pptx,
    "png": extract_image,
    "jpg": extract_image,
    "jpeg": extract_image,
    "gif": extract_image,
    "webp": extract_image,
    "bmp": extract_image,
}


def extract_by_ext(ext: str, data: bytes) -> tuple[str, bool]:
    """按扩展名分发抽取, 返回 (text, is_scanned)."""
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        raise ServiceNotReadyError(f"不支持的文件类型: {ext}")
    return extractor(data)  # type: ignore[operator]


def truncate(text: str, max_chars: int = MAX_CHARS) -> tuple[str, bool]:
    """超长截断, 返回 (text, truncated)."""
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars], True
